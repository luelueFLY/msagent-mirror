#!/usr/bin/python3
# -*- coding: utf-8 -*-
# -------------------------------------------------------------------------
# This file is part of the MindStudio project.
# Copyright (c) 2026 Huawei Technologies Co.,Ltd.
#
# MindStudio is licensed under Mulan PSL v2.
# You can use this software according to the terms and conditions of the Mulan PSL v2.
# You may obtain a copy of Mulan PSL v2 at:
#
#          http://license.coscl.org.cn/MulanPSL2
#
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND,
# EITHER EXPRESS OR IMPLIED, INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT,
# MERCHANTABILITY OR FIT FOR A PARTICULAR PURPOSE.
# See the Mulan PSL v2 for more details.
# -------------------------------------------------------------------------

"""compare_runs.py —— 多 run 对比（A/B 基线 vs 当前）

消费 msagent-profiler-breakdown 产出的统一拆解树，产出多 sheet 的 xlsx 对比表（阶段、稳态 step、
层间 wall_time、层内子模块占比）与一句话结论，服务于「改前 vs 改后」的量化差异判定。

用法：
  python compare_runs.py --base-dir <runA> --comp-dir <runB> \
                         [--labels "baseline,current"] [--output compare.xlsx]

- 每个目录内读取固定产物名 breakdown.json（统一拆解树；stage.json 为同内容别名；
  旧格式 stage.json + steps.json 兼容回退）、layers.json 可选。
- 依赖 openpyxl 写 xlsx；缺省（ImportError）时降级为逐 sheet 的 CSV + 控制台汇总。

设计依据：设计方案 §3.7 / §5（原 breakdown-compare 下沉），report-template.md「多 run 对比 xlsx」。
"""

import argparse
import csv
import json
import os

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - 离线无第三方依赖时降级
    HAS_OPENPYXL = False


def load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_json_opt(path):
    """可选读取：文件不存在（如 layer-decompose 暂缓无 layers.json）返回空 dict。"""
    if not os.path.isfile(path):
        return {}
    return load_json(path)


def load_run(rundir):
    """读取一个 run 的产物，返回规范化 dict。

    - 新格式：breakdown.json（统一拆解树，顶层端到端 -> children 阶段 + 单次执行分支；
      stage.json 为同内容别名）。
    - 旧格式（仅 stage.json + steps.json）兼容回退。
    layers 可选（layer-decompose 暂缓期间为空）。
    """
    tree_path = os.path.join(rundir, "breakdown.json")
    stage_path = os.path.join(rundir, "stage.json")
    if os.path.isfile(tree_path):
        tree = load_json(tree_path)
    else:
        tree = load_json(stage_path)
    children = tree.get("children", [])
    stages = [c for c in children if c.get("level") == "阶段"]
    step_node = next((c for c in children if c.get("level") == "单次执行"), None)
    steps = step_node or load_json_opt(os.path.join(rundir, "steps.json"))
    return {
        "dir": rundir,
        "stage": tree,
        "stages": stages,
        "steps": steps,
        "layers": load_json_opt(os.path.join(rundir, "layers.json")),
    }


def ms(ns):
    return ns / 1e6 if isinstance(ns, (int, float)) else None


def diff_pct(base, comp):
    """(comp - base) / base * 100；base 为 0/None 或 comp 为 None 返回 None。"""
    if base in (None, 0) or comp is None:
        return None
    return (comp - base) / base * 100.0


def fmt_ms(v, nd=2):
    x = ms(v)
    return ("%.*f" % (nd, x)) if x is not None else "-"


def _csv_cell(v):
    """CSV 单元格安全化：转字符串；以 = + - @ 开头（Excel 公式注入向量）时加单引号前缀。"""
    s = str(v)
    return "'" + s if s.startswith(("=", "+", "-", "@")) else s


def run_metrics(r):
    """抽取可对比指标。"""
    stage = r["stage"]
    stages = r["stages"]
    steps = r["steps"]
    layers = r["layers"]

    stage_map = {c["name"]: c.get("self_time", 0) for c in stages}
    steady = [s for s in steps.get("steps", []) if not s.get("is_warmup")
              and not s.get("is_prefill")]
    ss = steps.get("stats", {})

    best = min(steady, key=lambda s: s.get("wall_time", 0)) if steady else None
    worst = max(steady, key=lambda s: s.get("wall_time", 0)) if steady else None

    layer_walls = {layer["name"]: layer.get("wall_time", 0)
                   for layer in layers.get("layers", [])}
    layer_names = [layer["name"] for layer in layers.get("layers", [])]
    sub_pct = {s["name"]: s.get("pct", 0) for s in layers.get("submodule_totals", [])}

    return {
        "session_wall": stage.get("wall_time"),
        "stage_map": stage_map,
        "steady": steady,
        "stats": ss,
        "best_step": best,
        "worst_step": worst,
        "layer_walls": layer_walls,
        "layer_names": layer_names,
        "sub_pct": sub_pct,
        "num_layers": layers.get("num_layers"),
    }


def union_keys(a, b):
    keys = []
    for k in list(a.keys()) + list(b.keys()):
        if k not in keys:
            keys.append(k)
    return keys


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F2F2")


def _num_row(ws, label, a, b, fmt_a, fmt_b, pct_fmt="%.2f%%"):
    """追加一行数值对比：标量 a/b，展示时格式化，Δ% 用原始值计算。"""
    dp = diff_pct(a, b)
    ws.append([label, fmt_a(a), fmt_b(b),
               (pct_fmt % dp) if dp is not None else "-"])


def write_xlsx(ma, mb, base_label, comp_label, out_path):
    wb = Workbook()

    # Sheet 1 总览
    ws = wb.active
    ws.title = "总览"
    ws.append(["指标", base_label, comp_label, "Δ%"])
    _num_row(ws, "端到端 wall_time(ms)", ma["session_wall"], mb["session_wall"],
             lambda v: fmt_ms(v), lambda v: fmt_ms(v))
    sa, sb = ma["stats"], mb["stats"]
    _num_row(ws, "稳态 step P50(ms)", sa.get("p50"), sb.get("p50"),
             lambda v: fmt_ms(v), lambda v: fmt_ms(v))
    _num_row(ws, "稳态 step P99(ms)", sa.get("p99"), sb.get("p99"),
             lambda v: fmt_ms(v), lambda v: fmt_ms(v))
    _num_row(ws, "稳态 step mean(ms)", sa.get("mean"), sb.get("mean"),
             lambda v: fmt_ms(v), lambda v: fmt_ms(v))
    _num_row(ws, "稳态 step CV", sa.get("cv"), sb.get("cv"),
             lambda v: "%.4f" % (v or 0), lambda v: "%.4f" % (v or 0))
    ws.append(["层数", ma["num_layers"], mb["num_layers"], "-"])
    style_header(ws, 1, 4)

    # Sheet 2 阶段对比
    ws2 = wb.create_sheet("阶段对比")
    ws2.append(["阶段", base_label + "(ms)", comp_label + "(ms)", "Δ%"])
    for k in union_keys(ma["stage_map"], mb["stage_map"]):
        ws2.append([k, fmt_ms(ma["stage_map"].get(k)), fmt_ms(mb["stage_map"].get(k)),
                    ("%.2f%%" % diff_pct(ma["stage_map"].get(k), mb["stage_map"].get(k)))
                    if diff_pct(ma["stage_map"].get(k), mb["stage_map"].get(k)) is not None else "-"])
    style_header(ws2, 1, 4)

    # Sheet 3 稳态 step（best/worst + 序列对齐）
    ws3 = wb.create_sheet("稳态step")
    ws3.append(["指标", base_label + "(ms)", comp_label + "(ms)", "Δ%"])
    bb, bw = ma["best_step"], ma["worst_step"]
    cb, cw = mb["best_step"], mb["worst_step"]
    for label, a, b in (("最好 step", bb, cb), ("最差 step", bw, cw)):
        av = a.get("wall_time") if a else None
        bv = b.get("wall_time") if b else None
        _num_row(ws3, label, av, bv, lambda v: fmt_ms(v), lambda v: fmt_ms(v))
    ws3.append(["稳态 count", sa.get("count"), sb.get("count"), "-"])
    style_header(ws3, 1, 4)
    ws3.append([])
    ws3.append(["step", base_label + "(ms)", comp_label + "(ms)", "Δ%"])
    n = min(len(ma["steady"]), len(mb["steady"]))
    for i in range(n):
        a = ma["steady"][i].get("wall_time")
        b = mb["steady"][i].get("wall_time")
        dp = diff_pct(a, b)
        ws3.append(["#%d" % i, fmt_ms(a), fmt_ms(b),
                    ("%.2f%%" % dp) if dp is not None else "-"])

    # Sheet 4 层对比（层名取两侧并集，与阶段/子模块占比口径一致）
    ws4 = wb.create_sheet("层对比")
    ws4.append(["层", base_label + "(ms)", comp_label + "(ms)", "Δ%"])
    for k in union_keys(ma["layer_walls"], mb["layer_walls"]):
        a = ma["layer_walls"].get(k)
        b = mb["layer_walls"].get(k)
        dp = diff_pct(a, b)
        ws4.append([k, fmt_ms(a), fmt_ms(b), ("%.2f%%" % dp) if dp is not None else "-"])
    style_header(ws4, 1, 4)

    # Sheet 5 子模块占比
    ws5 = wb.create_sheet("子模块占比")
    ws5.append(["子模块", base_label + "(%)", comp_label + "(%)", "Δ(pp)"])
    for k in union_keys(ma["sub_pct"], mb["sub_pct"]):
        a = ma["sub_pct"].get(k)
        b = mb["sub_pct"].get(k)
        ws5.append([k, ("%.1f" % a) if a is not None else "-",
                    ("%.1f" % b) if b is not None else "-",
                    ("%.2f" % (b - a)) if (a is not None and b is not None) else "-"])
    style_header(ws5, 1, 4)

    wb.save(out_path)
    return out_path


def write_csv_fallback(ma, mb, base_label, comp_label, out_prefix):
    """openpyxl 缺失时降级：逐 sheet 写 CSV。"""
    sheets = {
        "overview": [["metric", base_label, comp_label],
                     ["session_wall_ms", fmt_ms(ma["session_wall"]), fmt_ms(mb["session_wall"])],
                     ["step_p50_ms", fmt_ms(ma["stats"].get("p50")), fmt_ms(mb["stats"].get("p50"))],
                     ["step_p99_ms", fmt_ms(ma["stats"].get("p99")), fmt_ms(mb["stats"].get("p99"))],
                     ["step_cv", "%.4f" % (ma["stats"].get("cv") or 0),
                      "%.4f" % (mb["stats"].get("cv") or 0)]],
        "stage": [["stage", base_label + "(ms)", comp_label + "(ms)"]] +
                 [[k, fmt_ms(ma["stage_map"].get(k)), fmt_ms(mb["stage_map"].get(k))]
                  for k in union_keys(ma["stage_map"], mb["stage_map"])],
        "layer": [["layer", base_label + "(ms)", comp_label + "(ms)"]] +
                 [[k, fmt_ms(ma["layer_walls"].get(k)), fmt_ms(mb["layer_walls"].get(k))]
                  for k in union_keys(ma["layer_walls"], mb["layer_walls"])],
        "submodule": [["submodule", base_label + "(%)", comp_label + "(%)"]] +
                     [[k, ma["sub_pct"].get(k), mb["sub_pct"].get(k)]
                      for k in union_keys(ma["sub_pct"], mb["sub_pct"])],
    }
    paths = []
    for name, rows in sheets.items():
        path = out_prefix + "_" + name + ".csv"
        with open(path, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            for row in rows:
                w.writerow([_csv_cell(c) for c in row])
        paths.append(path)
    return paths


def main(argv=None):
    p = argparse.ArgumentParser(description="多 run 对比（A/B）")
    p.add_argument("--base-dir", required=True, help="基线 run 目录（含 stage/steps/layers.json）")
    p.add_argument("--comp-dir", required=True, help="对比 run 目录（含 stage/steps/layers.json）")
    p.add_argument("--labels", default="baseline,current", help="两个 run 的标签，逗号分隔")
    p.add_argument("--output", default="compare.xlsx", help="输出 xlsx 路径（openpyxl 缺失时降级 CSV 前缀）")
    args = p.parse_args(argv)

    labels = [s.strip() for s in args.labels.split(",")]
    base_label = labels[0] if len(labels) > 0 else "baseline"
    comp_label = labels[1] if len(labels) > 1 else "current"

    base = load_run(args.base_dir)
    comp = load_run(args.comp_dir)
    ma = run_metrics(base)
    mb = run_metrics(comp)

    a_p50 = ms(ma["stats"].get("p50"))
    b_p50 = ms(mb["stats"].get("p50"))
    if a_p50 is not None and b_p50 is not None:
        d = diff_pct(ma["stats"]["p50"], mb["stats"]["p50"])
        trend = "变慢" if d > 0 else ("变快" if d < 0 else "持平")
        line = "稳态 step P50：%s=%.2fms → %s=%.2fms（%+.2f%%，%s）" % (
            base_label, a_p50, comp_label, b_p50, (d or 0), trend)
    else:
        line = "无稳态 step 数据可比对"

    if HAS_OPENPYXL:
        out = write_xlsx(ma, mb, base_label, comp_label, args.output)
        print("[OK] %s" % out)
    else:
        prefix = os.path.splitext(args.output)[0]
        paths = write_csv_fallback(ma, mb, base_label, comp_label, prefix)
        print("[WARN] openpyxl 未安装，降级为 CSV：%s" % ", ".join(paths))

    print("[INFO] " + line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
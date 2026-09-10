#!/usr/bin/python3
# -*- coding: utf-8 -*-
# -------------------------------------------------------------------------
# This file is part of the MindStudio project.
# Copyright (c) 2026 Huawei Technologies Co.,Ltd.
#
# MindStudio is licensed under Mulan PSL v2.
# You can use this software according to the terms and conditions of the Mulan PSL v2.
# You may obtain a copy of Mulan PSL v2 at:
#
#          http://license.coscl.org.cn/MulanPSL2
#
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND,
# EITHER EXPRESS OR IMPLIED, INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT,
# MERCHANTABILITY OR FIT FOR A PARTICULAR PURPOSE.
# See the Mulan PSL v2 for more details.
# -------------------------------------------------------------------------

"""
compare_runs.py 单元测试
"""

import json
import os
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO

sys.path.insert(
    0,
    os.path.join(
        os.path.dirname(__file__),
        "..",
        "..",
        "..",
        "..",
        "skills",
        "profiler",
        "msagent-profiler-breakdown",
        "scripts",
    ),
)

import compare_runs
from compare_runs import (
    HAS_OPENPYXL,
    _csv_cell,
    diff_pct,
    fmt_ms,
    load_json_opt,
    load_run,
    ms,
    run_metrics,
    union_keys,
    write_csv_fallback,
    write_xlsx,
)


def _sample_steps():
    return {
        "level": "单次执行",
        "name": "单次执行边界 + per-step 序列",
        "num_steps": 5,
        "num_steady": 3,
        "anchor": ["forward"],
        "steps": [
            {
                "index": 0,
                "name": "step.0",
                "start": 0,
                "end": 100,
                "wall_time": 100,
                "is_warmup": True,
                "warmup_reason": "graph_capture",
                "is_prefill": False,
            },
            {
                "index": 1,
                "name": "step.1",
                "start": 100,
                "end": 200,
                "wall_time": 100,
                "is_warmup": True,
                "warmup_reason": "compile",
                "is_prefill": False,
            },
            {
                "index": 2,
                "name": "step.2",
                "start": 200,
                "end": 300,
                "wall_time": 100,
                "is_warmup": False,
                "warmup_reason": None,
                "is_prefill": False,
            },
            {
                "index": 3,
                "name": "step.3",
                "start": 300,
                "end": 450,
                "wall_time": 150,
                "is_warmup": False,
                "warmup_reason": None,
                "is_prefill": False,
            },
            {
                "index": 4,
                "name": "step.4",
                "start": 450,
                "end": 500,
                "wall_time": 50,
                "is_warmup": False,
                "warmup_reason": None,
                "is_prefill": False,
            },
        ],
        "stats": {
            "count": 3,
            "mean": 100,
            "p50": 100,
            "p99": 150,
            "min": 50,
            "max": 150,
            "variance": 1666.67,
            "cv": 0.408,
        },
    }


def _sample_tree(framework="vllm"):
    return {
        "level": "端到端",
        "name": "端到端（会话）",
        "framework": framework,
        "task_type": "inference",
        "wall_time": 1000000000,
        "self_time": 800000000,
        "children": [
            {"level": "阶段", "name": "forward", "self_time": 500000000},
            {"level": "阶段", "name": "prepare_input", "self_time": 300000000},
            _sample_steps(),
        ],
    }


def _write_run(rundir, tree, steps_json=None, layers=None):
    """写一个 run 目录：breakdown.json + 可选 steps.json / layers.json。"""
    os.makedirs(rundir, exist_ok=True)
    with open(os.path.join(rundir, "breakdown.json"), "w", encoding="utf-8") as f:
        json.dump(tree, f, ensure_ascii=False)
    if steps_json is not None:
        with open(os.path.join(rundir, "steps.json"), "w", encoding="utf-8") as f:
            json.dump(steps_json, f, ensure_ascii=False)
    if layers is not None:
        with open(os.path.join(rundir, "layers.json"), "w", encoding="utf-8") as f:
            json.dump(layers, f, ensure_ascii=False)


class TestLoadJsonOpt(unittest.TestCase):
    def test_existing(self):
        tmpdir = tempfile.mkdtemp()
        try:
            path = os.path.join(tmpdir, "a.json")
            with open(path, "w", encoding="utf-8") as f:
                json.dump({"x": 1}, f)
            self.assertEqual(load_json_opt(path), {"x": 1})
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)

    def test_missing_returns_empty(self):
        tmpdir = tempfile.mkdtemp()
        try:
            self.assertEqual(load_json_opt(os.path.join(tmpdir, "nope.json")), {})
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)


class TestLoadRun(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_new_format(self):
        rundir = os.path.join(self.tmpdir, "run")
        _write_run(rundir, _sample_tree())
        r = load_run(rundir)
        self.assertEqual(len(r["stages"]), 2)
        self.assertEqual([s["name"] for s in r["stages"]], ["forward", "prepare_input"])
        self.assertEqual(r["steps"]["num_steady"], 3)
        self.assertEqual(r["layers"], {})

    def test_old_format_fallback(self):
        rundir = os.path.join(self.tmpdir, "run_old")
        tree = _sample_tree()
        tree["children"] = []
        _write_run(rundir, tree, steps_json=_sample_steps())
        # 模拟旧格式：只有 stage.json，无 breakdown.json
        os.rename(os.path.join(rundir, "breakdown.json"), os.path.join(rundir, "stage.json"))
        r = load_run(rundir)
        self.assertEqual(r["stages"], [])
        self.assertEqual(r["steps"]["num_steps"], 5)  # 回退读取 steps.json

    def test_missing_raises(self):
        rundir = os.path.join(self.tmpdir, "run_missing")
        os.makedirs(rundir, exist_ok=True)
        with self.assertRaises(FileNotFoundError):
            load_run(rundir)


class TestScalarHelpers(unittest.TestCase):
    def test_ms(self):
        self.assertEqual(ms(500000), 0.5)
        self.assertEqual(ms(1_000_000), 1.0)
        self.assertIsNone(ms(None))
        self.assertIsNone(ms("x"))

    def test_diff_pct(self):
        self.assertEqual(diff_pct(100, 150), 50.0)
        self.assertAlmostEqual(diff_pct(150, 100), 100 * (100 - 150) / 150)
        self.assertIsNone(diff_pct(0, 5))
        self.assertIsNone(diff_pct(None, 5))
        self.assertIsNone(diff_pct(5, None))

    def test_fmt_ms(self):
        self.assertEqual(fmt_ms(1000000), "1.00")
        self.assertEqual(fmt_ms(None), "-")
        self.assertEqual(fmt_ms(1500000), "1.50")

    def test_csv_cell_formula_injection(self):
        self.assertEqual(_csv_cell("=SUM(A1)"), "'=SUM(A1)")
        self.assertEqual(_csv_cell("+cmd"), "'+cmd")
        self.assertEqual(_csv_cell("@x"), "'@x")
        self.assertEqual(_csv_cell("normal"), "normal")


class TestUnionKeys(unittest.TestCase):
    def test_union_preserves_order(self):
        self.assertEqual(union_keys({"a": 1, "b": 2}, {"b": 3, "c": 4}), ["a", "b", "c"])

    def test_empty(self):
        self.assertEqual(union_keys({}, {}), [])


class TestRunMetrics(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.rundir = os.path.join(self.tmpdir, "run")
        layers = {
            "num_layers": 32,
            "layers": [
                {"name": "layer.0", "wall_time": 1000},
                {"name": "layer.1", "wall_time": 2000},
            ],
            "submodule_totals": [
                {"name": "attention", "pct": 60.0},
                {"name": "mlp", "pct": 40.0},
            ],
        }
        _write_run(self.rundir, _sample_tree(), layers=layers)

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_metrics(self):
        m = run_metrics(load_run(self.rundir))
        self.assertEqual(m["session_wall"], 1000000000)
        self.assertEqual(m["stage_map"]["forward"], 500000000)
        self.assertEqual(len(m["steady"]), 3)
        self.assertEqual(m["best_step"]["wall_time"], 50)
        self.assertEqual(m["worst_step"]["wall_time"], 150)
        self.assertEqual(m["num_layers"], 32)
        self.assertEqual(m["layer_walls"]["layer.1"], 2000)
        self.assertEqual(m["sub_pct"]["attention"], 60.0)

    def test_metrics_empty_steps(self):
        rundir = os.path.join(self.tmpdir, "run_empty")
        tree = _sample_tree()
        tree["children"] = [
            {"level": "阶段", "name": "forward", "self_time": 1},
        ]
        _write_run(rundir, tree)
        m = run_metrics(load_run(rundir))
        self.assertEqual(m["steady"], [])
        self.assertIsNone(m["best_step"])
        self.assertIsNone(m["worst_step"])


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl 未安装，跳过 xlsx 断言")
class TestWriteXlsx(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.ma = run_metrics(load_run(self._make_run("a", 100)))
        self.mb = run_metrics(load_run(self._make_run("b", 120)))

    def _make_run(self, name, p50_ns):
        rundir = os.path.join(self.tmpdir, name)
        tree = _sample_tree()
        tree["wall_time"] = p50_ns * 10
        tree["children"][2]["stats"]["p50"] = p50_ns
        _write_run(rundir, tree)
        return rundir

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_write_xlsx(self):
        out_path = os.path.join(self.tmpdir, "compare.xlsx")
        write_xlsx(self.ma, self.mb, "base", "comp", out_path)
        self.assertTrue(os.path.isfile(out_path))
        from openpyxl import load_workbook

        wb = load_workbook(out_path)
        self.assertEqual(wb.sheetnames, ["总览", "阶段对比", "稳态step", "层对比", "子模块占比"])


class TestWriteCsvFallback(unittest.TestCase):
    def _make_metrics(self, p50_ns):
        tmpdir = tempfile.mkdtemp()
        try:
            rundir = os.path.join(tmpdir, "run")
            tree = _sample_tree()
            tree["wall_time"] = p50_ns * 10
            tree["children"][2]["stats"]["p50"] = p50_ns
            _write_run(rundir, tree)
            return run_metrics(load_run(rundir))
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)

    def test_write_csv(self):
        tmpdir = tempfile.mkdtemp()
        try:
            ma = self._make_metrics(100_000)  # 会话 wall = 1.0 ms
            mb = self._make_metrics(120_000)  # 会话 wall = 1.2 ms
            paths = write_csv_fallback(ma, mb, "base", "comp", os.path.join(tmpdir, "cmp"))
            self.assertEqual(len(paths), 4)
            for name in ("overview", "stage", "layer", "submodule"):
                self.assertTrue(os.path.isfile(os.path.join(tmpdir, "cmp_%s.csv" % name)), name)
            with open(os.path.join(tmpdir, "cmp_overview.csv"), encoding="utf-8", newline="") as f:
                rows = f.read().splitlines()
            self.assertEqual(rows[0], "metric,base,comp")
            self.assertEqual(rows[1], "session_wall_ms,1.00,1.20")
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)

    def test_layer_sheet_includes_comp_only_layer(self):
        # 层对比取两侧并集：comp 独有 layer 也要出现在对比结果中
        tmpdir = tempfile.mkdtemp()
        try:

            def _metrics(rundir, names):
                layers = {
                    "num_layers": len(names),
                    "layers": [{"name": n, "wall_time": i + 1} for i, n in enumerate(names)],
                    "submodule_totals": [],
                }
                _write_run(rundir, _sample_tree(), layers=layers)
                return run_metrics(load_run(rundir))

            ma = _metrics(os.path.join(tmpdir, "a"), ["layer.0", "layer.1"])
            mb = _metrics(os.path.join(tmpdir, "b"), ["layer.0", "layer.2"])
            write_csv_fallback(ma, mb, "base", "comp", os.path.join(tmpdir, "cmp"))
            with open(os.path.join(tmpdir, "cmp_layer.csv"), encoding="utf-8", newline="") as f:
                rows = [r.split(",")[0] for r in f.read().splitlines()]
            self.assertIn("layer.0", rows)
            self.assertIn("layer.1", rows)
            self.assertIn("layer.2", rows)  # comp 独有层
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)


class TestMain(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.base = os.path.join(self.tmpdir, "base")
        self.comp = os.path.join(self.tmpdir, "comp")
        _write_run(self.base, _sample_tree())
        _write_run(self.comp, _sample_tree(framework="vllm"))

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_main(self):
        out_path = os.path.join(self.tmpdir, "compare.xlsx")
        out = StringIO()
        with redirect_stdout(out):
            rc = compare_runs.main(
                [
                    "--base-dir",
                    self.base,
                    "--comp-dir",
                    self.comp,
                    "--labels",
                    "baseline,current",
                    "--output",
                    out_path,
                ]
            )
        self.assertEqual(rc, 0)
        log = out.getvalue()
        self.assertIn("稳态 step P50", log)
        if HAS_OPENPYXL:
            self.assertTrue(os.path.isfile(out_path))
        else:
            self.assertTrue(os.path.isfile(os.path.join(self.tmpdir, "compare_overview.csv")))


if __name__ == "__main__":
    unittest.main()
